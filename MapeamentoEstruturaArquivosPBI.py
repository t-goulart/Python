import os  # Usado para varrer a pasta, montar caminhos e verificar existência de arquivos/diretórios
import json  # Usado para interpretar o Layout do PBIX (JSON) e extrair os nomes das páginas
import zipfile  # Usado para abrir o .pbix, que internamente é um arquivo ZIP
import datetime  # Usado para obter data/hora atuais (registro no Excel de histórico)
import openpyxl  # Usado para criar/atualizar a planilha Excel de histórico de mapeamento
from openpyxl.worksheet.table import Table, TableStyleInfo  # Usado para criar/expandir a tabela nomeada do Excel de histórico
from openpyxl.styles import Alignment  # Usado para desativar a quebra de texto automática nas células do Excel

# Tenta carregar o pbixray
try:  # Tenta importar a biblioteca que lê o modelo compilado de dentro do .pbix
    from pbixray import PBIXRay  # CORREÇÃO: a classe se chama "PBIXRay" — o código anterior importava "PBIX", que não existe na biblioteca
    varPbixrayDisponivel = True  # Sinaliza que a leitura completa do modelo está disponível
except Exception:  # A biblioteca não está instalada nesta máquina ou falhou ao carregar
    varPbixrayDisponivel = False  # Sinaliza que só será possível extrair as páginas do relatório

r"""
Objetivo: varrer a pasta de painéis em produção, ler a estrutura interna de cada
arquivo .pbix (tabelas e colunas, medidas DAX, colunas e tabelas calculadas,
consultas Power Query/M, relacionamentos e páginas do relatório) e gerar um arquivo
.txt individual por painel, registrando o resultado de cada mapeamento (sucesso ou
falha) em uma base Excel histórica.

Quais são os dados que precisamos?
- Acesso de leitura à pasta de rede dos painéis em produção (varPastaOrigem).
- Biblioteca "pbixray" instalada na máquina que executa o script, pois o modelo de
  dados do .pbix é compilado em formato binário e não pode ser lido apenas com
  zipfile/json. Instalação: pip install pbixray

POR QUE O RELATÓRIO ESTAVA SAINDO QUASE VAZIO?
A versão anterior gerava um .txt apenas com "Nenhuma tabela explícita encontrada",
"Nenhuma medida DAX encontrada" e "Nenhuma consulta M encontrada". Três defeitos
somados causavam isso:

1) IMPORT INCORRETO DA BIBLIOTECA: o código fazia "from pbixray import PBIX", mas a
   classe exposta pela biblioteca se chama "PBIXRay". O import sempre levantava
   ImportError, o flag de disponibilidade ficava False e a leitura completa do
   modelo NUNCA era executada — o script caía direto no fallback manual.

2) CAMINHOS INTERNOS ERRADOS NO FALLBACK: o fallback procurava os arquivos "Layout"
   e "DataModelSchema" na raiz do ZIP. No .pbix real, o layout fica em
   "Report/Layout" (com a subpasta), e "DataModelSchema" simplesmente NÃO EXISTE em
   arquivos .pbix — ele só existe em arquivos .pbit (template). Como nenhuma das
   duas condições era satisfeita, o fallback devolvia listas vazias e o relatório
   saía sem conteúdo.

3) MODELO DO .PBIX É BINÁRIO, NÃO JSON: em um .pbix o modelo fica no fluxo
   "DataModel", que é um pacote binário compilado (xVelocity/Analysis Services), e
   não um JSON legível. Tentar interpretá-lo com json.loads jamais funcionaria — é
   exatamente por isso que a leitura do modelo depende da biblioteca pbixray.

Além disso, mesmo que o import funcionasse, os nomes de propriedades e colunas
usados estavam desatualizados em relação à API da biblioteca:
  - "model.measures" -> o correto é "model.dax_measures" (coluna "Name", e não "MeasureName")
  - "model.queries"  -> o correto é "model.power_query"  (coluna "TableName", e não "QueryName")
  - "model.tables"   -> devolve apenas a LISTA de nomes de tabelas; o detalhamento de
                        colunas e tipos está em "model.schema" (TableName/ColumnName/PandasDataType)

Qual é a ordem que os processos são executados?
1 - processarVarreduraPbi: lista os arquivos .pbix da pasta de origem (apenas o nível
    raiz, sem subpastas, como no comportamento original).
2 - Para cada painel (Fase 1), mapearPbix monta o relatório de texto com: páginas do
    relatório, tabelas e colunas, medidas DAX, colunas calculadas, tabelas calculadas,
    consultas Power Query (M) e relacionamentos.
3 - registrarHistoricoExcel grava uma linha na base Excel de histórico para cada
    painel processado, com o status "Sucesso" ou "Falha".
4 - Painéis que falharam na Fase 1 são reunidos em uma fila e reprocessados uma
    segunda vez na Fase 2, dando uma segunda chance a arquivos que só estavam
    temporariamente indisponíveis (em uso/salvando no momento da varredura).

Pontos de atenção:
- As páginas do relatório são lidas do fluxo "Report/Layout" (JSON em utf-16-le) e
  aparecem no relatório mesmo quando o pbixray não está instalado — é a única parte
  da estrutura que não depende da biblioteca.
- Se o pbixray não estiver instalado, o painel é marcado como "Falha" no Excel e o
  .txt explica exatamente o motivo, em vez de gerar um relatório vazio sem explicação.
- Tabelas automáticas de inteligência temporal criadas pelo próprio Power BI
  ("DateTableTemplate_..." e "LocalDateTable_...") aparecem no relatório como
  qualquer outra tabela, pois fazem parte real do modelo.
"""

# ==========================================
# CONFIGURAÇÕES DE CAMINHOS
# ==========================================
varPastaOrigem = r""  # Pasta onde ficam os painéis em produção
varPastaDestinoTxt = r""  # Pasta onde cada relatório .txt individual é salvo
varPastaDestinoExcel = r""  # Pasta onde fica a base Excel de histórico de mapeamento
varCaminhoExcel = os.path.join(varPastaDestinoExcel, "Nome do Arquivo de Monitoramento.xlsx")  # Caminho completo do arquivo Excel de histórico


def extrairPaginasDoLayout(varCaminhoPbix):
    """
    Objetivo: extrair os nomes das páginas (abas) do relatório Power BI, lendo o
    fluxo "Report/Layout" de dentro do .pbix.

    Entrada:
        varCaminhoPbix (str) -- caminho completo do arquivo .pbix a inspecionar.

    Retorno:
        list[str] com o nome de cada página do relatório, na ordem em que aparecem.
        Lista vazia se o layout não puder ser lido — esta parte nunca interrompe o
        mapeamento, pois é complementar às informações do modelo de dados.

    Ponto de atenção: o caminho correto dentro do ZIP é "Report/Layout" (com a
    subpasta) e o conteúdo vem codificado em utf-16-le. A versão anterior procurava
    apenas "Layout" na raiz, que nunca existe, e por isso nunca listava as páginas.
    """
    varListaPaginas = []  # Lista que vai acumular os nomes das páginas encontradas

    try:  # Tenta abrir o .pbix como ZIP e ler o layout do relatório
        with zipfile.ZipFile(varCaminhoPbix, "r") as varArquivoZip:  # Abre o .pbix, que internamente é um pacote ZIP
            varNomesInternos = varArquivoZip.namelist()  # Lista todos os fluxos/arquivos existentes dentro do pacote

            varCaminhoLayout = "Report/Layout" if "Report/Layout" in varNomesInternos else None  # Localiza o layout no caminho correto (com subpasta)

            if varCaminhoLayout is None:  # Alguns pacotes antigos podem trazer o layout na raiz
                varCaminhoLayout = "Layout" if "Layout" in varNomesInternos else None  # Tenta o caminho alternativo, por compatibilidade

            if varCaminhoLayout is not None:  # Só prossegue se algum dos caminhos de layout existir
                varConteudoLayout = varArquivoZip.read(varCaminhoLayout)  # Lê os bytes brutos do layout
                varLayoutJson = json.loads(varConteudoLayout.decode("utf-16-le", errors="ignore"))  # Interpreta o JSON, que vem codificado em utf-16-le

                for varSecao in varLayoutJson.get("sections", []):  # Percorre cada seção do layout, que corresponde a uma página do relatório
                    varListaPaginas.append(varSecao.get("displayName", varSecao.get("name")))  # Usa o nome exibido da página; se não houver, usa o nome interno
    except Exception:  # Qualquer falha na leitura do layout é ignorada silenciosamente
        pass  # As páginas são informação complementar — sua ausência não invalida o restante do relatório

    return varListaPaginas  # Devolve os nomes das páginas encontradas


def mapearPbix(varCaminhoPbix, varCaminhoSaida):
    """
    Objetivo: ler a estrutura completa de um painel Power BI (.pbix) e gravar o
    relatório correspondente em um arquivo .txt.

    Entrada:
        varCaminhoPbix (str) -- caminho completo do arquivo .pbix a ser mapeado.
        varCaminhoSaida (str) -- caminho completo do arquivo .txt de relatório a gerar.

    Retorno:
        bool -- True se o modelo de dados foi lido com sucesso; False se a leitura
        do modelo falhou (pbixray ausente, arquivo em uso, painel corrompido etc.),
        caso em que o painel entra na fila de repescagem da Fase 2.
    """
    varConteudoTxt = []  # Lista que acumula, linha a linha, o conteúdo do relatório de texto
    varConteudoTxt.append("==================================================")  # Abre o cabeçalho visual do relatório
    varConteudoTxt.append("RELATÓRIO DE ESTRUTURA DO PAINEL POWER BI (.PBIX)")  # Título do relatório
    varConteudoTxt.append(f"Arquivo: {varCaminhoPbix}")  # Identifica qual painel este relatório documenta
    varConteudoTxt.append("==================================================\n")  # Fecha o cabeçalho visual do relatório

    varSucesso = False  # Indica se a leitura do modelo de dados foi concluída com sucesso

    # PÁGINAS DO RELATÓRIO (independe do pbixray — lidas direto do Report/Layout)
    varListaPaginas = extrairPaginasDoLayout(varCaminhoPbix)  # Extrai os nomes das páginas do relatório
    varConteudoTxt.append("📄 PÁGINAS DO PAINEL:")  # Abre a seção de páginas do relatório
    if varListaPaginas:  # Verifica se alguma página foi encontrada
        for varNomePagina in varListaPaginas:  # Percorre cada página encontrada
            varConteudoTxt.append(f"   • {varNomePagina}")  # Registra o nome da página no relatório
    else:  # Nenhuma página pôde ser lida do layout
        varConteudoTxt.append("   Nenhuma página encontrada no layout do relatório.")  # Informa a ausência de páginas no relatório
    varConteudoTxt.append("")  # Linha em branco separando a seção de páginas das próximas

    if not varPbixrayDisponivel:  # A biblioteca que lê o modelo compilado não está instalada nesta máquina
        varConteudoTxt.append("❌ Não foi possível ler o modelo de dados deste painel.")  # Informa o problema no próprio relatório
        varConteudoTxt.append("   Motivo: a biblioteca 'pbixray' não está instalada nesta máquina.")  # Explica a causa exata da limitação
        varConteudoTxt.append("   Solução: execute 'pip install pbixray' no Python que roda esta automação.")  # Orienta a correção
        varConteudoTxt.append("   (O modelo de dados do .pbix é binário compilado e não pode ser lido apenas com zipfile/json.)")  # Esclarece por que a biblioteca é obrigatória

        os.makedirs(os.path.dirname(varCaminhoSaida), exist_ok=True)  # Garante que a pasta de destino do relatório existe antes de salvar
        with open(varCaminhoSaida, "w", encoding="utf-8") as varArquivoSaida:  # Abre o arquivo de saída para escrita em UTF-8
            varArquivoSaida.write("\n".join(varConteudoTxt))  # Grava o relatório parcial (apenas páginas + explicação da limitação)

        return False  # Sinaliza falha: sem a biblioteca, o mapeamento do modelo não pôde ser feito

    try:  # Tenta ler o modelo de dados completo do painel via pbixray
        varModelo = PBIXRay(varCaminhoPbix)  # Abre o painel e carrega os metadados do modelo compilado

        # 1. TABELAS E COLUNAS
        varConteudoTxt.append("\n--- [ 1. TABELAS E TIPOS DE DADOS ] ---\n")  # Abre a seção de tabelas do relatório
        varSchema = varModelo.schema  # DataFrame com uma linha por coluna do modelo (TableName, ColumnName, PandasDataType)
        if not varSchema.empty:  # Verifica se o modelo tem ao menos uma coluna mapeada
            for varNomeTabela in varSchema["TableName"].unique():  # Percorre cada tabela distinta presente no modelo
                varConteudoTxt.append(f"📌 Tabela: {varNomeTabela}")  # Registra o nome da tabela no relatório
                varColunasTabela = varSchema[varSchema["TableName"] == varNomeTabela]  # Filtra apenas as colunas pertencentes a esta tabela
                varConteudoTxt.append("   Campos / Colunas:")  # Introduz a listagem de colunas da tabela
                for _, varLinhaColuna in varColunasTabela.iterrows():  # Percorre cada coluna desta tabela
                    varConteudoTxt.append(f"     • {varLinhaColuna.get('ColumnName')}: {varLinhaColuna.get('PandasDataType')}")  # Registra o nome e o tipo de dado da coluna
                varConteudoTxt.append("")  # Linha em branco separando esta tabela da próxima
        else:  # O modelo não trouxe nenhuma coluna
            varConteudoTxt.append("Nenhuma tabela explícita encontrada.")  # Informa a ausência de tabelas no relatório

        # 2. MEDIDAS DAX
        varConteudoTxt.append("\n--- [ 2. MEDIDAS DAX ] ---\n")  # Abre a seção de medidas do relatório
        varMedidas = varModelo.dax_measures  # DataFrame com as medidas DAX (TableName, Name, Expression, DisplayFolder, Description)
        if not varMedidas.empty:  # Verifica se o modelo tem ao menos uma medida
            for _, varLinhaMedida in varMedidas.iterrows():  # Percorre cada medida do modelo
                varConteudoTxt.append(f"📐 Medida [{varLinhaMedida.get('TableName')}]: {varLinhaMedida.get('Name')} =")  # Registra a tabela de origem e o nome da medida
                varConteudoTxt.append(f"   {varLinhaMedida.get('Expression')}")  # Registra a expressão DAX completa da medida
                varConteudoTxt.append("-" * 40)  # Linha separadora visual entre medidas
        else:  # O modelo não tem medidas cadastradas
            varConteudoTxt.append("Nenhuma medida DAX encontrada.")  # Informa a ausência de medidas no relatório

        # 3. COLUNAS CALCULADAS (DAX)
        varConteudoTxt.append("\n--- [ 3. COLUNAS CALCULADAS (DAX) ] ---\n")  # Abre a seção de colunas calculadas do relatório
        varColunasCalculadas = varModelo.dax_columns  # DataFrame com as colunas calculadas em DAX (TableName, ColumnName, Expression)
        if not varColunasCalculadas.empty:  # Verifica se o modelo tem ao menos uma coluna calculada
            for _, varLinhaCalculada in varColunasCalculadas.iterrows():  # Percorre cada coluna calculada do modelo
                varConteudoTxt.append(f"🧮 Coluna [{varLinhaCalculada.get('TableName')}]: {varLinhaCalculada.get('ColumnName')} =")  # Registra a tabela de origem e o nome da coluna calculada
                varConteudoTxt.append(f"   {varLinhaCalculada.get('Expression')}")  # Registra a expressão DAX que define a coluna
                varConteudoTxt.append("-" * 40)  # Linha separadora visual entre colunas calculadas
        else:  # O modelo não tem colunas calculadas
            varConteudoTxt.append("Nenhuma coluna calculada encontrada.")  # Informa a ausência de colunas calculadas no relatório

        # 4. TABELAS CALCULADAS (DAX)
        varConteudoTxt.append("\n--- [ 4. TABELAS CALCULADAS (DAX) ] ---\n")  # Abre a seção de tabelas calculadas do relatório
        varTabelasCalculadas = varModelo.dax_tables  # DataFrame com as tabelas criadas por expressão DAX (TableName, Expression)
        if not varTabelasCalculadas.empty:  # Verifica se o modelo tem ao menos uma tabela calculada
            for _, varLinhaTabelaCalc in varTabelasCalculadas.iterrows():  # Percorre cada tabela calculada do modelo
                varConteudoTxt.append(f"🧾 Tabela Calculada: {varLinhaTabelaCalc.get('TableName')} =")  # Registra o nome da tabela calculada
                varConteudoTxt.append(f"   {varLinhaTabelaCalc.get('Expression')}")  # Registra a expressão DAX que gera a tabela
                varConteudoTxt.append("-" * 40)  # Linha separadora visual entre tabelas calculadas
        else:  # O modelo não tem tabelas calculadas
            varConteudoTxt.append("Nenhuma tabela calculada encontrada.")  # Informa a ausência de tabelas calculadas no relatório

        # 5. FONTES DE DADOS E POWER QUERY (M)
        varConteudoTxt.append("\n\n--- [ 5. FONTES DE DADOS E POWER QUERY (M) ] ---\n")  # Abre a seção de consultas Power Query do relatório
        varConsultasM = varModelo.power_query  # DataFrame com as consultas Power Query do painel (TableName, Expression)
        if not varConsultasM.empty:  # Verifica se o painel tem ao menos uma consulta M
            for _, varLinhaConsulta in varConsultasM.iterrows():  # Percorre cada consulta Power Query do painel
                varConteudoTxt.append(f"🔍 Consulta: {varLinhaConsulta.get('TableName')}")  # Registra o nome da consulta (que dá nome à tabela carregada)
                varConteudoTxt.append(f"Código M:\n{varLinhaConsulta.get('Expression')}")  # Registra o código M completo da consulta
                varConteudoTxt.append("=" * 40)  # Linha separadora visual entre consultas
        else:  # O painel não tem consultas M (ex: conexão direta/live connection)
            varConteudoTxt.append("Nenhuma consulta M encontrada.")  # Informa a ausência de consultas M no relatório

        # 6. RELACIONAMENTOS
        varConteudoTxt.append("\n--- [ 6. RELACIONAMENTOS ENTRE TABELAS ] ---\n")  # Abre a seção de relacionamentos do relatório
        varRelacionamentos = varModelo.relationships  # DataFrame com os relacionamentos do modelo (origem, destino, cardinalidade etc.)
        if not varRelacionamentos.empty:  # Verifica se o modelo tem ao menos um relacionamento
            for _, varLinhaRelacao in varRelacionamentos.iterrows():  # Percorre cada relacionamento do modelo
                varConteudoTxt.append(f"🔗 {dict(varLinhaRelacao)}")  # Registra todos os atributos do relacionamento, preservando o que a biblioteca devolver
        else:  # O modelo não tem relacionamentos (ou a biblioteca não conseguiu lê-los)
            varConteudoTxt.append("Nenhum relacionamento encontrado.")  # Informa a ausência de relacionamentos no relatório

        varSucesso = True  # Marca que a leitura do modelo foi concluída com sucesso

    except Exception as varErroLeitura:  # Captura qualquer falha na leitura do modelo (arquivo em uso, painel corrompido, formato não suportado)
        varConteudoTxt.append(f"\n❌ Erro ao ler metadados: {varErroLeitura}")  # Registra o erro no próprio relatório, para facilitar o diagnóstico
        varSucesso = False  # Marca a falha, para que o painel entre na fila de repescagem

    # Grava o arquivo TXT
    os.makedirs(os.path.dirname(varCaminhoSaida), exist_ok=True)  # Garante que a pasta de destino do relatório existe antes de salvar
    with open(varCaminhoSaida, "w", encoding="utf-8") as varArquivoSaida:  # Abre o arquivo de saída para escrita em UTF-8
        varArquivoSaida.write("\n".join(varConteudoTxt))  # Grava todo o conteúdo acumulado, uma linha por item da lista

    return varSucesso  # Devolve o resultado do mapeamento deste painel


def registrarHistoricoExcel(varCaminhoExcelPath, varNomePbi, varCaminhoRede, varStatus):
    """
    Objetivo: criar (na primeira execução) ou atualizar (nas execuções seguintes) a
    base Excel de histórico de mapeamento, adicionando uma linha por painel
    processado, dentro de uma tabela nomeada do Excel, sem quebra de texto nas
    células.

    Entrada:
        varCaminhoExcelPath (str) -- caminho completo do arquivo Excel de histórico.
        varNomePbi (str) -- nome do painel (sem extensão) que foi processado.
        varCaminhoRede (str) -- caminho completo do arquivo .pbix na rede.
        varStatus (str) -- resultado do mapeamento ("Sucesso" ou "Falha").

    Retorno:
        Nenhum. A função grava o arquivo Excel atualizado diretamente em disco.
    """
    os.makedirs(os.path.dirname(varCaminhoExcelPath), exist_ok=True)  # Garante que a pasta de destino do Excel existe antes de salvar

    varAgora = datetime.datetime.now()  # Marca o instante exato deste registro de histórico
    varDataStr = varAgora.strftime("%d/%m/%Y")  # Formata a data no padrão brasileiro dd/mm/aaaa
    varHoraStr = varAgora.strftime("%H:%M")  # Formata a hora no padrão hh:mm

    varCabecalhos = ["Nome", "Local", "Data Mapeamento", "Hora", "Status"]  # Define os títulos das colunas da base de histórico

    if not os.path.exists(varCaminhoExcelPath):  # Verifica se a base de histórico ainda não existe (primeira execução)
        varWorkbook = openpyxl.Workbook()  # Cria um novo arquivo Excel em memória
        varPlanilha = varWorkbook.active  # Recupera a aba criada por padrão
        varPlanilha.title = "Mapeamento"  # Nomeia a aba de dados
        varPlanilha.append(varCabecalhos)  # Escreve a linha de cabeçalho na primeira linha da planilha
        varPlanilha.append([varNomePbi, varCaminhoRede, varDataStr, varHoraStr, varStatus])  # Escreve a primeira linha de dados (este mapeamento)

        varTabela = Table(displayName="tMapeamentoEstruturaPaineis", ref="A1:E2")  # Cria a tabela nomeada cobrindo cabeçalho + primeira linha de dados
        varEstiloTabela = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)  # Define o estilo visual padrão da tabela
        varTabela.tableStyleInfo = varEstiloTabela  # Aplica o estilo visual à tabela recém-criada
        varPlanilha.add_table(varTabela)  # Registra a tabela na planilha
    else:  # A base de histórico já existe: abre e apenas empilha uma nova linha
        varWorkbook = openpyxl.load_workbook(varCaminhoExcelPath)  # Abre o arquivo existente, preservando todo o histórico já registrado
        varPlanilha = varWorkbook.active  # Recupera a aba de dados já existente

        varPlanilha.append([varNomePbi, varCaminhoRede, varDataStr, varHoraStr, varStatus])  # Empilha a nova linha de histórico abaixo das já existentes

        if "tMapeamentoEstruturaPaineis" in varPlanilha.tables:  # Verifica se a tabela nomeada já existe nesta planilha
            varTabela = varPlanilha.tables["tMapeamentoEstruturaPaineis"]  # Recupera a referência da tabela existente
            varTabela.ref = f"A1:E{varPlanilha.max_row}"  # Expande o range da tabela para cobrir a nova linha recém-adicionada

    varSemQuebraTexto = Alignment(wrap_text=False)  # Define o alinhamento sem quebra automática de texto, para manter cada linha compacta
    for varLinha in varPlanilha.iter_rows(min_row=1, max_row=varPlanilha.max_row, min_col=1, max_col=5):  # Percorre todas as linhas (cabeçalho + dados) das 5 colunas da base
        for varCelula in varLinha:  # Percorre cada célula da linha atual
            varCelula.alignment = varSemQuebraTexto  # Aplica o alinhamento sem quebra de texto a esta célula

    varWorkbook.save(varCaminhoExcelPath)  # Grava o arquivo Excel atualizado em disco


def processarVarreduraPbi():
    """
    Objetivo: orquestrar toda a varredura — listar os painéis .pbix da pasta de
    origem, mapear cada um (Fase 1), registrar o resultado no Excel de histórico, e
    dar uma segunda chance (Fase 2) aos painéis que falharam.

    Entrada:
        Nenhuma. Utiliza as constantes definidas na seção de configurações.

    Retorno:
        int -- quantidade total de painéis mapeados com sucesso (somando Fase 1 e
        Fase 2).
    """
    os.makedirs(varPastaDestinoTxt, exist_ok=True)  # Garante que a pasta de destino dos relatórios .txt existe antes de iniciar

    varPaineisPendentes = []  # Lista que acumula, na Fase 1, os painéis que falharam e precisam de uma segunda tentativa
    varTotalSucesso = 0  # Contador do total de painéis mapeados com sucesso, somando as duas fases

    if not varPbixrayDisponivel:  # Avisa logo no início se a biblioteca essencial não está instalada
        print("⚠️ Atenção: a biblioteca 'pbixray' não está instalada — os relatórios sairão sem o modelo de dados.")  # Alerta o operador antes mesmo de começar
        print("   Instale com: pip install pbixray\n")  # Orienta a correção imediata

    print(f"🔍 [FASE 1] Iniciando varredura na pasta (SEM SUBPASTAS): {varPastaOrigem}\n")  # Informa no terminal que a Fase 1 está começando

    for varNomeArquivo in os.listdir(varPastaOrigem):  # Percorre os itens da pasta de origem (apenas o nível raiz, sem subpastas)
        varCaminhoCompleto = os.path.join(varPastaOrigem, varNomeArquivo)  # Monta o caminho completo do item atual

        if os.path.isfile(varCaminhoCompleto) and varNomeArquivo.lower().endswith(".pbix") and not varNomeArquivo.startswith("~$"):  # Considera apenas arquivos .pbix, ignorando arquivos temporários de lock (~$)
            varNomeSemExtensao = os.path.splitext(varNomeArquivo)[0]  # Extrai o nome do painel sem a extensão .pbix
            varCaminhoTxt = os.path.join(varPastaDestinoTxt, f"{varNomeSemExtensao}.txt")  # Monta o caminho do relatório .txt correspondente a este painel

            print(f"⚡ Mapeando Painel: {varNomeArquivo}...")  # Informa no terminal qual painel está sendo mapeado agora

            varSucesso = mapearPbix(varCaminhoCompleto, varCaminhoTxt)  # Executa o mapeamento completo deste painel

            if varSucesso:  # Verifica se o mapeamento deste painel foi bem-sucedido
                registrarHistoricoExcel(varCaminhoExcel, varNomeSemExtensao, varCaminhoCompleto, "Sucesso")  # Registra o sucesso no Excel de histórico
                varTotalSucesso += 1  # Soma mais um mapeamento bem-sucedido ao total geral
                print(f"   └─ ✅ Sucesso!")  # Confirma no terminal que este painel foi mapeado com sucesso
            else:  # O mapeamento deste painel falhou
                registrarHistoricoExcel(varCaminhoExcel, varNomeSemExtensao, varCaminhoCompleto, "Falha")  # Registra a falha no Excel de histórico
                varPaineisPendentes.append((varNomeSemExtensao, varCaminhoCompleto, varCaminhoTxt))  # Adiciona este painel à fila de repescagem da Fase 2
                print(f"   └─ ⚠️ Erro ao processar arquivo")  # Informa no terminal que este painel entrou na fila de repescagem

    if varPaineisPendentes:  # Verifica se algum painel ficou pendente na Fase 1
        print(f"\n🔄 [FASE 2] Retornando para tentar mapear {len(varPaineisPendentes)} painel(eis) pendente(s)...\n")  # Informa no terminal que a Fase 2 está começando

        for varNomeSemExtensao, varCaminhoPbix, varCaminhoTxt in varPaineisPendentes:  # Percorre cada painel que ficou pendente na Fase 1
            print(f"⚡ Segunda tentativa: {os.path.basename(varCaminhoPbix)}...")  # Informa no terminal qual painel está sendo remapeado agora
            varSucesso = mapearPbix(varCaminhoPbix, varCaminhoTxt)  # Executa uma nova tentativa de mapeamento completo deste painel

            if varSucesso:  # Verifica se a segunda tentativa foi bem-sucedida
                registrarHistoricoExcel(varCaminhoExcel, varNomeSemExtensao, varCaminhoPbix, "Sucesso")  # Registra o sucesso da repescagem no Excel de histórico
                varTotalSucesso += 1  # Soma mais um mapeamento bem-sucedido ao total geral
                print(f"   └─ ✅ Sucesso na repescagem!")  # Confirma no terminal que este painel foi mapeado com sucesso na segunda tentativa
            else:  # A segunda tentativa também falhou
                print(f"   └─ ❌ Falha persistente no arquivo.")  # Informa no terminal que este painel continua com problema

    print(f"\n🎉 Processo concluído! Total de painéis Power BI mapeados com sucesso: {varTotalSucesso}")  # Resumo final da varredura completa, com o total de sucessos

    return varTotalSucesso  # Devolve a quantidade total de sucessos


def main():
    """
    Objetivo: orquestrar a execução — chamar processarVarreduraPbi() e reportar no
    terminal qualquer erro inesperado ocorrido durante o processamento.

    Entrada:
        Nenhuma.

    Retorno:
        Nenhum.
    """
    try:  # Envolve todo o processo para capturar qualquer erro inesperado
        processarVarreduraPbi()  # Executa a varredura completa dos painéis
    except Exception as varErro:  # Captura qualquer falha inesperada ocorrida durante o processamento
        print(f"❌ Erro: {varErro}")  # Exibe a descrição do erro no terminal


if __name__ == "__main__":  # Garante que o processo só roda quando o arquivo é executado diretamente
    main()  # Executa a automação completa
